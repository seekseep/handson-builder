#!/usr/bin/env python3
"""
各レクチャーの example/data/ に置く練習用ブック（.xlsx）を生成する。

マクロ付きのブック（.xlsm）は作らない。VBA は受講者が自分で貼り付けて書くもので、
出所の分からないマクロ付きブックを開かせる教材にしたくないため。
.xlsx は Google スプレッドシートにもそのままアップロードできる。

使い方（保守者のみ。受講者は実行しない）:
    pip install openpyxl
    python3 tools/build_workbooks.py

データは乱数を使うなら種を固定し、何度実行しても同じ中身になるようにする。
セル番地はコード（vba/*.bas・gas/*.gs）と対応しているので、レイアウトを変えたら両方を直す。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
SECTIONS = ROOT / "sections"

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ROWS = [
    # 単価, 数量
    (120, 5),
    (180, 5),
    (250, 5),
    (550, 4),
    (680, 1),
]


def write_sample_sheet(ws):
    for col, h in enumerate(["単価", "数量", "金額"], start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BOX
    for r, (price, quantity) in enumerate(ROWS, start=2):
        for col, v in enumerate([price, quantity, None], start=1):
            ws.cell(row=r, column=col, value=v).border = BOX


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"[build_workbooks] {path.relative_to(ROOT)}")


def main():
    # 同じブックを使うレクチャーを並べる（節ごとに ZIP で配るので、各 example/data/ に置く）
    for lecture in ("01-basics/01-hello",):
        wb = Workbook()
        ws = wb.active
        ws.title = "シート1"
        write_sample_sheet(ws)
        save(wb, SECTIONS / lecture / "example" / "data" / "sample.xlsx")


if __name__ == "__main__":
    main()
